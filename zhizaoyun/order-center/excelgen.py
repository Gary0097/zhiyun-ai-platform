# -*- coding: utf-8 -*-
"""订单 Excel 生成（给财务）：订单信息 / 物料明细 / 解析说明 三张表。"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITLE_FILL = PatternFill("solid", fgColor="1E3A8A")
HEAD_FILL = PatternFill("solid", fgColor="DBEAFE")
K_FILL = PatternFill("solid", fgColor="F1F5F9")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _title(ws, text, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(size=14, bold=True, color="FFFFFF")
    c.fill = TITLE_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30


def _head(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="1E3A8A")
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")


def build(rec):
    """rec: db.get_order 返回的字典（含 parsed_json / adapted_json）。"""
    parsed = rec.get("parsed_json") or {}
    adapted = rec.get("adapted_json") or {}
    prod = adapted.get("匹配产品") or {}
    materials = adapted.get("物料清单") or []
    warnings = rec.get("warnings") or []

    wb = Workbook()

    # 明细合计（订单信息页成本以此为准，避免与物料明细表口径不一致）
    mat_total = 0.0
    for m in materials:
        qty = m.get("数量") or 0
        price = m.get("标准单价") or 0
        amt = m.get("金额")
        if amt in (None, ""):
            try:
                amt = float(qty) * float(price)
            except Exception:  # noqa: BLE001
                amt = 0
        mat_total += float(amt or 0)
    cost_show = round(mat_total, 2) if materials else (
        adapted.get("预估物料成本") if adapted.get("预估物料成本") is not None else rec.get("cost") or 0)

    # ---- Sheet1 订单信息 ----
    ws = wb.active
    ws.title = "订单信息"
    _title(ws, "金汉隆智造云 · 客户订单标准化工单", 4)
    rows = [
        ("销售订单编号(SO)", rec.get("so_no") or "", "标准化工单号(MO)", rec.get("mo_no") or ""),
        ("客户名称", parsed.get("客户名称") or rec.get("customer") or "", "客户类型", parsed.get("客户类型") or ""),
        ("客户联系人", parsed.get("客户姓名") or "", "联系方式", parsed.get("联系方式") or ""),
        ("品名", parsed.get("品名") or "", "数量 / 单位", "%s %s" % (parsed.get("数量") if parsed.get("数量") is not None else "",
                                                              parsed.get("单位") or "")),
        ("设备型号", prod.get("产品名称") or parsed.get("设备型号建议") or "", "产品族", prod.get("产品族") or ""),
        ("模板分类", adapted.get("模板分类") or rec.get("template") or "", "是否自产设备", parsed.get("是否自产设备") or ""),
        ("交期", parsed.get("交期") or rec.get("due_date") or "", "付款方式", parsed.get("付款方式") or ""),
        ("收货地址", parsed.get("收货地址") or "", "原始单据来源", parsed.get("原始单据来源") or rec.get("source_type") or ""),
        ("质保金", parsed.get("质保金") or "", "有无押模", parsed.get("有无押模") or ""),
        ("预估物料成本(元)", cost_show,
         "制单时间", (rec.get("created_at") or "")[:19].replace("T", " ")),
    ]
    r = 3
    for k1, v1, k2, v2 in rows:
        for col, (k, v) in enumerate(((k1, v1), (k2, v2))):
            kc = ws.cell(row=r, column=1 + col * 2, value=k)
            vc = ws.cell(row=r, column=2 + col * 2, value=v)
            kc.font = Font(bold=True)
            kc.fill = K_FILL
            kc.border = BORDER
            vc.border = BORDER
            vc.alignment = Alignment(wrap_text=True, vertical="center")
        r += 1
    for col, w in zip("ABCD", (22, 42, 20, 42)):
        ws.column_dimensions[col].width = w

    # ---- Sheet2 物料明细 ----
    ws2 = wb.create_sheet("物料明细")
    headers = ["序号", "物料编号", "物料名称", "物料类别", "单位", "数量", "标准单价(元)", "金额(元)", "备注"]
    _title(ws2, "物料明细（依据知识库物料主数据匹配）", len(headers))
    _head(ws2, 3, headers)
    r = 4
    total = 0.0
    for i, m in enumerate(materials, start=1):
        qty = m.get("数量") or 0
        price = m.get("标准单价") or 0
        amt = m.get("金额")
        if amt in (None, ""):
            try:
                amt = float(qty) * float(price)
            except Exception:  # noqa: BLE001
                amt = 0
        total += float(amt or 0)
        vals = [i, m.get("物料编号") or "", m.get("物料名称") or "", m.get("物料类别") or "",
                m.get("单位") or "", qty, price, round(float(amt or 0), 2), m.get("备注") or ""]
        for c, v in enumerate(vals, start=1):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.border = BORDER
        r += 1
    if not materials:
        ws2.cell(row=r, column=1, value="（智能体未返回物料清单，请检查解析结果或人工补充）")
        r += 1
    else:
        ws2.cell(row=r, column=3, value="合计").font = Font(bold=True)
        tc = ws2.cell(row=r, column=8, value=round(total, 2))
        tc.font = Font(bold=True)
        tc.fill = K_FILL
    for i, w in enumerate((6, 14, 24, 12, 8, 10, 14, 14, 26), start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---- Sheet3 解析说明 ----
    ws3 = wb.create_sheet("解析说明")
    _title(ws3, "AI 解析与匹配说明（供财务复核）", 4)
    _head(ws3, 3, ["类型", "内容", "", ""])
    r = 4

    def block(title, items):
        nonlocal r
        c = ws3.cell(row=r, column=1, value=title)
        c.font = Font(bold=True, color="1E3A8A")
        c.fill = K_FILL
        c.border = BORDER
        ws3.cell(row=r, column=2, value="").border = BORDER
        r += 1
        if not items:
            items = ["（无）"]
        for it in items:
            ws3.cell(row=r, column=1, value="").border = BORDER
            cell = ws3.cell(row=r, column=2, value=str(it))
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1

    block("解析说明", [parsed.get("解析说明") or "（无）"])
    block("判类依据", [adapted.get("判类依据") or "（无）"])
    fmap = adapted.get("字段映射") or []
    block("字段映射", ["%s → %s：%s" % (m.get("订单原文", ""), m.get("标准字段", ""), m.get("值", "")) for m in fmap])
    block("待人工确认", (parsed.get("待人工确认") or []) + (adapted.get("待人工确认") or []) + list(warnings))
    block("核对提示", adapted.get("核对提示") or [])
    block("原始单据文本", [(rec.get("raw_text") or "")[:3000]])
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def filename(rec):
    cust = (rec.get("customer") or (rec.get("parsed_json") or {}).get("客户名称") or "客户")
    short = str(cust)[:12].replace("/", "_").replace("\\", "_")
    return "订单_%s_%s.xlsx" % (short, rec.get("so_no") or datetime.now().strftime("%Y%m%d%H%M%S"))
