# -*- coding: utf-8 -*-
"""金汉隆 · 销售分析平台 — 销售BI / 客户价值 / 业绩归因 + 数据导入智能分析"""
import csv
import io
import json
import os
import sqlite3
from pathlib import Path

import requests
from flask import Flask, g, jsonify, render_template, request, send_file

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("SALES_DB", BASE_DIR / "data" / "sales.db"))

# ---- 多企业演示：按访问 IP(Host 头) 切换企业库；未配置时行为不变 ----
import sys as _sys
_kit = os.environ.get("ENTERPRISE_KIT")
if _kit and os.path.isdir(_kit) and _kit not in _sys.path:
    _sys.path.insert(0, _kit)
try:
    import enterprise_selector as _ent
except Exception:
    _ent = None


def ent_db_path():
    if _ent is not None:
        p = _ent.db_path("sales.db")
        if p:
            return Path(p)
    return DB_PATH


LLM_BASE = os.environ.get("SALES_LLM_BASE", "http://api.computing.lzwcai.com:3110/v1")
LLM_KEY = os.environ.get("SALES_LLM_KEY", "")
LLM_MODEL = os.environ.get("SALES_LLM_MODEL", "qwen3.8-flash")

app = Flask(__name__)


if _ent is not None:
    @app.before_request
    def _ent_set_host():
        _ent.set_host(request.host)


def db():
    if "db" not in g:
        g.db = sqlite3.connect(ent_db_path())
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    conn = g.pop("db", None)
    if conn is not None:
        conn.close()


# ----------------------------------------------------------------------------
# 页面
# ----------------------------------------------------------------------------
@app.route("/")
def home():
    return render_template("home.html")


@app.route("/bi")
def bi_page():
    return render_template("bi.html")


@app.route("/crm")
def crm_page():
    return render_template("crm.html")


@app.route("/commission")
def commission_page():
    return render_template("commission.html")


@app.route("/data")
def data_page():
    return render_template("data.html")


# ----------------------------------------------------------------------------
# 驾驶舱 / BI 智能分析
# ----------------------------------------------------------------------------
@app.route("/api/overview")
def api_overview():
    d = db()
    op = d.execute("SELECT COUNT(*) n, SUM(amount) sales, SUM(profit) profit "
                   "FROM order_profit").fetchone()
    customers = d.execute("SELECT COUNT(DISTINCT customer) c FROM order_profit").fetchone()["c"]
    warns = d.execute("SELECT COUNT(*) c FROM risk_warn").fetchone()["c"]
    tags = d.execute("SELECT tag, COUNT(*) n, SUM(total_amount) v FROM customer_value "
                     "GROUP BY tag").fetchall()
    margin = (op["profit"] or 0) / op["sales"] * 100 if op["sales"] else 0
    return jsonify({
        "order_count": op["n"], "total_sales": op["sales"] or 0,
        "total_profit": op["profit"] or 0, "margin": round(margin, 2),
        "customer_count": customers, "warn_count": warns,
        "tags": [dict(r) for r in tags],
    })


@app.route("/api/bi/products")
def api_bi_products():
    rows = db().execute(
        "SELECT product, sales, profit, COALESCE(CASE WHEN sales>0 AND profit IS NOT NULL "
        "THEN profit/sales*100 END, 0) margin "
        "FROM product_profit ORDER BY sales DESC").fetchall()
    return jsonify({"rows": [dict(r) for r in rows]})


@app.route("/api/bi/margin-trend")
def api_bi_margin_trend():
    """月度毛利率走势：合同号解析签约月（YYYYMM），月度聚合销售额/毛利。

    毛利为测算值（真实成本数据未接入，由 simulate_margin.py 批次生成）。
    可选参数 from/to（YYYYMM），默认 202401 ~ 当前月。
    """
    import re as _re

    def _ym(v):
        s = (v or "").strip()
        return s if _re.fullmatch(r"20\d{4}", s) else None

    today = "".join(str(__import__("datetime").date.today()).split("-")[:2])
    ym_from = _ym(request.args.get("from")) or "202401"
    ym_to = _ym(request.args.get("to")) or today
    agg = {}
    for r in db().execute("SELECT contract_no, amount, profit FROM order_profit").fetchall():
        m = _re.search(r"(20\d{6})", r["contract_no"] or "")
        if not m:
            continue
        k = m.group(1)
        try:
            if not (1 <= int(k[4:6]) <= 12 and 1 <= int(k[6:8]) <= 31):
                continue
        except ValueError:
            continue
        if not (ym_from <= k[:6] <= ym_to):
            continue
        a = agg.setdefault(k[:6], [0.0, 0.0])
        a[0] += r["amount"] or 0
        a[1] += r["profit"] or 0
    rows = [{"month": k[:4] + "-" + k[4:], "sales": round(v[0], 2),
             "profit": round(v[1], 2),
             "margin": round(v[1] / v[0] * 100, 2) if v[0] else 0}
            for k, v in sorted(agg.items()) if v[0] > 0]
    return jsonify({"rows": rows, "from": ym_from, "to": ym_to,
                    "note": "毛利为测算值"})


@app.route("/api/bi/reps")
def api_bi_reps():
    """销售人员业绩 + 复购率（基于订单-客户-销售映射计算）"""
    d = db()
    base = {r["rep"]: dict(r) for r in d.execute("SELECT * FROM sales_reps")}
    stat = {}
    for r in d.execute(
            "SELECT a.main_rep rep, o.customer, COUNT(*) n FROM attribution a "
            "JOIN order_profit o ON o.contract_no = a.contract_no "
            "GROUP BY a.main_rep, o.customer"):
        s = stat.setdefault(r["rep"], {"orders": 0, "customers": 0, "repeat_cust": 0})
        s["orders"] += r["n"]
        s["customers"] += 1
        if r["n"] > 1:
            s["repeat_cust"] += 1
    avg = (sum(s["repeat_cust"] / max(s["customers"], 1) for s in stat.values()) / len(stat)
           if stat else 0)
    avg_pct = avg * 100
    out = []
    for rep, s in stat.items():
        b = base.get(rep, {})
        rate = s["repeat_cust"] / max(s["customers"], 1) * 100
        out.append({"rep": rep, "orders": s["orders"], "customers": s["customers"],
                    "repeat_cust": s["repeat_cust"], "repeat_rate": round(rate, 1),
                    "sales": b.get("sales", 0), "est_profit": b.get("est_profit", 0),
                    "low_repeat": rate < avg_pct})
    out.sort(key=lambda x: -(x["sales"] or 0))
    return jsonify({"rows": out, "avg_repeat_rate": round(avg_pct, 1)})


@app.route("/api/bi/high-margin")
def api_bi_high_margin():
    """高毛利客户清单：按客户聚合，毛利率高于整体均值且销售额前列"""
    min_sales = request.args.get("min_sales", 100000, type=float)
    rows = db().execute(
        "SELECT customer, COUNT(*) orders, SUM(amount) sales, SUM(profit) profit, "
        "CASE WHEN SUM(amount)>0 THEN SUM(profit)/SUM(amount)*100 ELSE 0 END margin "
        "FROM order_profit GROUP BY customer HAVING SUM(amount) >= ? "
        "ORDER BY margin DESC, sales DESC LIMIT 30", (min_sales,)).fetchall()
    overall = db().execute(
        "SELECT CASE WHEN SUM(amount)>0 THEN SUM(profit)/SUM(amount)*100 ELSE 0 END m "
        "FROM order_profit").fetchone()["m"] or 0
    rows = [dict(r) for r in rows if r["margin"] and r["margin"] > overall]
    return jsonify({"rows": rows, "overall_margin": round(overall, 2), "note": "利润/毛利待真实BOM成本接入后启用"})


@app.route("/api/bi/customers")
def api_bi_customers():
    rows = db().execute(
        "SELECT customer, COUNT(*) orders, SUM(amount) sales, SUM(profit) profit, "
        "CASE WHEN SUM(amount)>0 THEN SUM(profit)/SUM(amount)*100 ELSE 0 END margin "
        "FROM order_profit GROUP BY customer ORDER BY sales DESC LIMIT 50").fetchall()
    return jsonify({"rows": [dict(r) for r in rows]})


@app.route("/api/warns")
def api_warns():
    rows = db().execute("SELECT * FROM risk_warn ORDER BY amount DESC LIMIT 100").fetchall()
    return jsonify({"rows": [dict(r) for r in rows]})


@app.route("/api/insight", methods=["POST"])
def api_insight():
    """智能解读：把聚合指标喂给大模型，输出分析结论"""
    if not LLM_KEY:
        return jsonify({"error": "未配置分析模型"}), 503
    d = db()
    ov = db().execute("SELECT COUNT(*) n, SUM(amount) s, SUM(profit) p FROM order_profit").fetchone()
    facts = {
        "订单数": ov["n"], "销售总额": round(ov["s"] or 0), "毛利": round(ov["p"] or 0),
        "产品Top5": [{"产品": r["product"], "销售额": round(r["sales"]),
                      "毛利率": round(r["profit"] / r["sales"] * 100, 1) if r["sales"] else 0}
                     for r in d.execute("SELECT * FROM product_profit ORDER BY sales DESC LIMIT 5")],
        "客户标签分布": {r["tag"]: r["n"] for r in d.execute(
            "SELECT tag, COUNT(*) n FROM customer_value GROUP BY tag")},
        "风险预警客户数": d.execute("SELECT COUNT(*) c FROM risk_warn").fetchone()["c"],
    }
    body = {
        "model": LLM_MODEL,
        "messages": [{"role": "user", "content":
            "你是销售分析专家。基于以下销售数据指标，输出4条以内的简明分析结论"
            "（每条一句话，聚焦毛利结构、客户价值、风险点，不要建议空话）：\n" +
            json.dumps(facts, ensure_ascii=False)}],
        "max_tokens": 300,
    }
    try:
        r = requests.post(LLM_BASE + "/chat/completions",
                          headers={"Authorization": "Bearer " + LLM_KEY},
                          json=body, timeout=60)
        r.raise_for_status()
        content = r.json()["choices"][0]["message"]["content"]
        return jsonify({"insight": content.strip()})
    except Exception as e:  # noqa: BLE001
        return jsonify({"error": f"分析模型调用失败：{e}"}), 502


# ----------------------------------------------------------------------------
# 客户价值（智能客户关系管理）
# ----------------------------------------------------------------------------
@app.route("/api/crm/list")
def api_crm_list():
    tag = request.args.get("tag", "")
    q = request.args.get("q", "").strip()
    where, args = "1=1", []
    if tag:
        where += " AND tag = ?"
        args.append(tag)
    if q:
        where += " AND customer LIKE ?"
        args.append(f"%{q}%")
    rows = db().execute(
        f"SELECT * FROM customer_value WHERE {where} ORDER BY total_amount DESC LIMIT 300",
        args).fetchall()
    return jsonify({"rows": [dict(r) for r in rows]})


@app.route("/api/crm/scatter")
def api_crm_scatter():
    rows = db().execute(
        "SELECT customer, total_amount, times, complaints, resp_min, score, tag "
        "FROM customer_value").fetchall()
    return jsonify({"rows": [dict(r) for r in rows]})


# ----------------------------------------------------------------------------
# 业绩归因（销售业绩智能统计 / 分佣报告）
# ----------------------------------------------------------------------------
@app.route("/api/commission/orders")
def api_commission_orders():
    page = max(int(request.args.get("page", 1)), 1)
    size = min(int(request.args.get("size", 20)), 100)
    q = request.args.get("q", "").strip()
    where, args = "", []
    if q:
        where = ("WHERE a.contract_no LIKE ? OR a.customer LIKE ? OR "
                 "a.main_rep LIKE ? OR a.tracker LIKE ?")
        args = [f"%{q}%"] * 4
    total = db().execute(
        f"SELECT COUNT(*) c FROM attribution a {where}", args).fetchone()["c"]
    rows = db().execute(
        f"SELECT a.*, o.amount FROM attribution a LEFT JOIN order_profit o "
        f"ON o.contract_no = a.contract_no {where} "
        f"ORDER BY o.amount DESC LIMIT ? OFFSET ?",
        args + [size, (page - 1) * size]).fetchall()
    return jsonify({"total": total, "rows": [dict(r) for r in rows]})


@app.route("/api/commission/report")
def api_commission_report():
    """分佣报告：主销售贡献 / 跟单贡献 / 协同贡献 三本账"""
    d = db()
    rep = {r["rep"]: {"role": "主销售", "orders": 0, "amount": 0.0}
           for r in d.execute("SELECT DISTINCT main_rep rep FROM attribution")}
    for r in d.execute("SELECT main_rep rep, COUNT(*) n, SUM(rep_amount) v "
                       "FROM attribution GROUP BY main_rep"):
        rep.setdefault(r["rep"], {"role": "主销售"})["orders"] = r["n"]
        rep[r["rep"]]["amount"] = r["v"] or 0
        rep[r["rep"]]["role"] = "主销售"
    for r in d.execute("SELECT tracker rep, COUNT(*) n, SUM(tracker_amount) v "
                       "FROM attribution GROUP BY tracker"):
        e = rep.setdefault(r["rep"], {"role": "跟单", "orders": 0, "amount": 0.0})
        e["role"] = "跟单" if e["role"] == "主销售" and not e["orders"] else \
            (e["role"] + "+跟单" if "跟单" not in e["role"] else e["role"])
        e["orders"] = e.get("orders", 0) + r["n"]
        e["amount"] = e.get("amount", 0) + (r["v"] or 0)
    for r in d.execute("SELECT coworker rep, COUNT(*) n, SUM(coworker_amount) v "
                       "FROM attribution WHERE coworker != '' GROUP BY coworker"):
        e = rep.setdefault(r["rep"], {"role": "协同", "orders": 0, "amount": 0.0})
        if "协同" not in e["role"]:
            e["role"] += "+协同"
        e["orders"] = e.get("orders", 0) + r["n"]
        e["amount"] = e.get("amount", 0) + (r["v"] or 0)
    rows = sorted(rep.items(), key=lambda kv: -kv[1]["amount"])
    total = db().execute(
        "SELECT SUM(rep_amount)+SUM(tracker_amount)+SUM(coworker_amount) t "
        "FROM attribution").fetchone()["t"]
    return jsonify({"rows": [{"name": k, **v} for k, v in rows], "total": total or 0})


@app.route("/api/commission/export")
def api_commission_export():
    rep = json.loads(api_commission_report().get_data(as_text=True))
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["\ufeff姓名", "角色", "订单数", "贡献金额"])
    for r in rep["rows"]:
        w.writerow([r["name"], r["role"], r["orders"], round(r["amount"], 2)])
    return send_file(io.BytesIO(buf.getvalue().encode("utf-8")), as_attachment=True,
                     download_name="销售分佣报告.csv", mimetype="text/csv")


# ----------------------------------------------------------------------------
# 数据管理：导入 / 导出 / 模板
# ----------------------------------------------------------------------------
DATA_TABLES = {
    "order_profit": {
        "name": "订单利润分析", "icon": "📦", "table": "order_profit",
        "headers": ["合同编号", "客户", "设备型号", "销售金额", "BOM采购成本", "订单毛利", "毛利率"],
        "cols": ["contract_no", "customer", "device", "amount", "bom_cost", "profit", "margin"],
        "numeric": {"amount", "bom_cost", "profit", "margin"},
    },
    "customer_value": {
        "name": "客户价值评估", "icon": "💎", "table": "customer_value",
        "headers": ["客户", "累计采购额", "采购次数", "投诉次数", "平均响应分钟", "价值评分", "客户标签"],
        "cols": ["customer", "total_amount", "times", "complaints", "resp_min", "score", "tag"],
        "numeric": {"total_amount", "resp_min", "score"},
    },
    "attribution": {
        "name": "销售贡献归因", "icon": "🤝", "table": "attribution",
        "headers": ["合同编号", "客户", "主销售", "跟单", "协同人员", "销售比例", "跟单比例",
                    "协同比例", "销售贡献金额", "跟单贡献金额", "协同贡献金额"],
        "cols": ["contract_no", "customer", "main_rep", "tracker", "coworker",
                 "rep_ratio", "tracker_ratio", "coworker_ratio",
                 "rep_amount", "tracker_amount", "coworker_amount"],
        "numeric": {"rep_ratio", "tracker_ratio", "coworker_ratio",
                    "rep_amount", "tracker_amount", "coworker_amount"},
    },
    "sales_reps": {
        "name": "销售人员业绩", "icon": "👤", "table": "sales_reps",
        "headers": ["销售人员", "订单数量", "销售额", "客户数量", "预计贡献毛利"],
        "cols": ["rep", "order_count", "sales", "customer_count", "est_profit"],
        "numeric": {"order_count", "sales", "customer_count", "est_profit"},
    },
    "risk_warn": {
        "name": "客户风险预警", "icon": "⚠️", "table": "risk_warn",
        "headers": ["客户", "采购额", "采购次数", "风险原因", "建议动作"],
        "cols": ["customer", "amount", "times", "reason", "action"],
        "numeric": {"amount", "times"},
    },
}


def _read_sheet(file_storage):
    filename = (file_storage.filename or "").lower()
    if filename.endswith(".csv"):
        stream = io.StringIO(file_storage.read().decode("utf-8-sig"))
        return list(csv.reader(stream))
    import openpyxl
    wb = openpyxl.load_workbook(file_storage, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return [[("" if c is None else str(c).strip()) for c in r]
            for r in ws.iter_rows(values_only=True)]


def _xlsx_response(spec, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec["name"]
    ws.append(spec["headers"])
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{spec['name']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/api/data/tables")
def api_data_tables():
    out = []
    for key, spec in DATA_TABLES.items():
        n = db().execute(f"SELECT COUNT(*) c FROM {spec['table']}").fetchone()["c"]
        out.append({"key": key, "name": spec["name"], "icon": spec["icon"],
                    "headers": spec["headers"], "count": n})
    return jsonify({"tables": out})


@app.route("/api/data/rows/<key>")
def api_data_rows(key):
    """表数据预览（只读，前 N 条，默认 50、上限 200）。"""
    spec = DATA_TABLES.get(key)
    if not spec:
        return jsonify({"error": "未知数据表"}), 404
    try:
        limit = min(max(int(request.args.get("limit", 50)), 1), 200)
    except ValueError:
        limit = 50
    rows = db().execute(
        f"SELECT {', '.join(spec['cols'])} FROM {spec['table']} LIMIT ?", (limit,)
    ).fetchall()
    return jsonify({"headers": spec["headers"],
                    "rows": [["" if v is None else v for v in tuple(r)] for r in rows]})


@app.route("/api/data/export/<key>")
def api_data_export(key):
    spec = DATA_TABLES.get(key)
    if not spec:
        return jsonify({"error": "未知数据表"}), 404
    rows = db().execute(f"SELECT {', '.join(spec['cols'])} FROM {spec['table']}").fetchall()
    return _xlsx_response(spec, [tuple(r) for r in rows])


@app.route("/api/data/template/<key>")
def api_data_template(key):
    spec = DATA_TABLES.get(key)
    if not spec:
        return jsonify({"error": "未知数据表"}), 404
    return _xlsx_response(spec, [])


@app.route("/api/data/import/<key>", methods=["POST"])
def api_data_import(key):
    spec = DATA_TABLES.get(key)
    if not spec:
        return jsonify({"error": "未知数据表"}), 404
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "缺少文件"}), 400
    try:
        data = _read_sheet(f)
        headers, raw_rows = data[0], data[1:]
        idx = {h.strip(): i for i, h in enumerate(headers)}
        missing = [h for h in spec["headers"] if h not in idx]
        if missing:
            raise ValueError(f"缺少列：{'、'.join(missing)}（请使用模板）")
        tuples = []
        for r in raw_rows:
            if not any(str(x).strip() for x in r):
                continue
            tup = []
            for col, header in zip(spec["cols"], spec["headers"]):
                v = r[idx[header]]
                if col in spec["numeric"]:
                    try:
                        v = float(str(v).replace(",", "") or 0)
                    except ValueError:
                        v = 0
                tup.append(v)
            tuples.append(tup)
    except Exception as e:  # noqa: BLE001
        return jsonify({"error": f"解析失败：{e}"}), 400
    if not tuples:
        return jsonify({"error": "没有可导入的数据行"}), 400
    conn = db()
    placeholders = ",".join("?" * len(spec["cols"]))
    cols = ",".join(spec["cols"])
    try:
        conn.execute(f"DELETE FROM {spec['table']}")
        conn.executemany(f"INSERT INTO {spec['table']} ({cols}) VALUES ({placeholders})", tuples)
        conn.commit()
    except sqlite3.Error as e:
        conn.rollback()
        return jsonify({"error": f"写入失败：{e}"}), 500
    return jsonify({"imported": len(tuples), "table": spec["name"]})


# ----------------------------------------------------------------------------
# AI 智能解析智能体（侧边聊天）：模型可发起只读 SQL 查询真实数据后作答
# ----------------------------------------------------------------------------
import re as _re
import time as _time

_SQL_OK = _re.compile(r"(?is)^\s*(select|with)\b")
_THINK_RE = _re.compile(r"(?is)<think>.*?</think>")


def _agent_schema():
    d = sqlite3.connect(f"file:{ent_db_path()}?mode=ro", uri=True)
    try:
        parts = []
        for (t,) in d.execute("SELECT name FROM sqlite_master WHERE type='table'"):
            cols = []
            for r in d.execute(f"PRAGMA table_info({t})"):
                col = r[1]
                nn = d.execute(
                    f"SELECT COUNT({col}) FROM {t}").fetchone()[0]
                cols.append(col + ("" if nn else "（全空,勿查询）"))
            n = d.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
            parts.append(f"{t}({', '.join(cols)}) -- {n}行")
        return "\n".join(parts)
    finally:
        d.close()


def _run_readonly_sql(sql):
    sql = sql.strip().rstrip(";")
    if not _SQL_OK.match(sql) or ";" in sql:
        raise ValueError("仅支持单条只读 SELECT/WITH 查询")
    if not _re.search(r"(?i)\blimit\b", sql):
        sql = f"SELECT * FROM ({sql}) LIMIT 30"
    d = sqlite3.connect(f"file:{ent_db_path()}?mode=ro", uri=True)
    d.row_factory = sqlite3.Row
    try:
        d.execute("PRAGMA query_only=ON")
        rows = [dict(r) for r in d.execute(sql).fetchall()]
        return {"columns": list(rows[0].keys()) if rows else [],
                "rows": rows[:30], "row_count": len(rows)}
    finally:
        d.close()


def _llm_chat(messages, max_tokens=900):
    body = {"model": LLM_MODEL, "messages": messages,
            "max_tokens": max_tokens, "temperature": 0.3}
    r = requests.post(LLM_BASE + "/chat/completions",
                      headers={"Authorization": "Bearer " + LLM_KEY},
                      json=body, timeout=90)
    r.raise_for_status()
    msg = r.json()["choices"][0]["message"]
    content = msg.get("content") or ""
    return _THINK_RE.sub("", content).strip()


# ----------------------------------------------------------------------------
# 伙伴云(Huoban) CRM 数据同步（业绩归因接通销售数据）
# ----------------------------------------------------------------------------
@app.route("/api/huoban/status")
def api_huoban_status():
    import glob

    d = db()
    d.execute("CREATE TABLE IF NOT EXISTS huoban_sync_log("
              "id INTEGER PRIMARY KEY AUTOINCREMENT, source TEXT, rows_read INTEGER, "
              "matched INTEGER, inserted INTEGER, synced_at TEXT, note TEXT)")
    last = d.execute("SELECT * FROM huoban_sync_log ORDER BY id DESC LIMIT 1").fetchone()
    files = sorted(os.path.basename(p) for p in
                   glob.glob(str(BASE_DIR / "data" / "huoban" / "*.json")))
    filled = d.execute("SELECT COUNT(*) c FROM attribution "
                       "WHERE rep_amount IS NOT NULL").fetchone()["c"]
    total = d.execute("SELECT COUNT(*) c FROM attribution").fetchone()["c"]
    return jsonify({"last": dict(last) if last else None, "files": files,
                    "attribution_total": total, "attribution_filled": filled})


@app.route("/api/huoban/sync", methods=["POST"])
def api_huoban_sync():
    import huoban_sync

    out = huoban_sync.run(str(DB_PATH), str(BASE_DIR / "data" / "huoban"))
    return jsonify(out)


@app.route("/api/agent/chat", methods=["POST"])
def api_agent_chat():
    """侧边聊天智能体：多轮对话，模型需要数据时返回 {"sql":...} 由服务端
    只读执行后回喂，最终基于真实数据回答。"""
    if not LLM_KEY:
        return jsonify({"error": "未配置分析模型"}), 503
    body = request.get_json(silent=True) or {}
    history = [(m.get("role"), m.get("content", ""))
               for m in (body.get("messages") or [])][-12:]
    history = [(r, c) for r, c in history if r in ("user", "assistant") and c]
    if not history:
        return jsonify({"error": "缺少对话内容"}), 400
    page = body.get("page") or ""
    pages = {"/bi": "销售BI分析", "/crm": "客户价值管理",
             "/commission": "业绩贡献归因", "/data": "数据管理", "/": "驾驶舱"}
    system = (
        "你是「金汉隆销售分析智能体」，嵌在销售分析平台的侧边聊天栏，用中文简明回答。\n"
        "数据库表结构（SQLite，只读）：\n" + _agent_schema() + "\n"
        "规则：\n"
        "1. 涉及数据的问题必须先查库：只输出一个 JSON 代码块 {\"sql\": \"SELECT ...\"}，"
        "只能用 SELECT/WITH，单条语句，记得 LIMIT。系统会把查询结果回给你。\n"
        "2. 拿到数据后直接给结论：数字准确、可分点、不超300字；金额用万元表述。\n"
        "3. 表中没有日期字段，不要按时间维度臆测；毛利率=profit/sales 或 margin 列。\n"
        "4. 客户/人名在不同表中的写法可能不一致：按名称过滤时一律用 LIKE '%关键词%' 模糊匹配，"
        "并在结论中注明数据来源表。\n"
        "5. 无关数据的问题（闲聊/常识）正常回答即可。\n"
        f"用户当前所在页面：{pages.get(page, page or '未知')}\n"
    )
    msgs = [{"role": "system", "content": system}] + [
        {"role": r, "content": c} for r, c in history]
    queries = []
    try:
        for _ in range(4):
            answer = _llm_chat(msgs)
            m = _re.search(r"(?s)```(?:json)?\s*(\{.*?\})\s*```|\{\s*\"sql\"\s*:\s*\"([^\"]+)\"",
                           answer)
            sql = None
            if m:
                try:
                    sql = (json.loads(m.group(1)).get("sql")
                           if m.group(1) else m.group(2))
                except Exception:  # noqa: BLE001
                    sql = None
            if not sql:
                return jsonify({"answer": answer, "queries": queries})
            try:
                result = _run_readonly_sql(sql)
            except Exception as e:  # noqa: BLE001
                result = {"error": f"SQL执行失败：{e}"}
            queries.append(sql)
            msgs.append({"role": "assistant", "content": answer})
            msgs.append({"role": "user", "content":
                         "查询结果(JSON)：" + json.dumps(result, ensure_ascii=False)[:6000]
                         + "\n请基于以上真实数据给出最终结论（不要再输出SQL）。"})
        return jsonify({"answer": "数据查询轮次过多，请换个更具体的问题。",
                        "queries": queries})
    except requests.RequestException as e:
        return jsonify({"error": f"分析模型调用失败：{e}"}), 502
    except Exception as e:  # noqa: BLE001
        return jsonify({"error": f"智能体异常：{e}"}), 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5060)), debug=False)
