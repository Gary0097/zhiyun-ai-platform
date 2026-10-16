# -*- coding: utf-8 -*-
"""金汉隆销售分析平台 — 数据导入：xlsx -> SQLite"""
import os
import sqlite3
import sys

import openpyxl

DB_PATH = os.environ.get("SALES_DB", os.path.join(os.path.dirname(__file__), "data", "sales.db"))


def f(v, default=0.0):
    try:
        if v is None or v == "":
            return default
        return float(str(v).replace(",", "").strip())
    except (ValueError, TypeError):
        return default


def s(v):
    return "" if v is None else str(v).strip()


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "data", "model.xlsx")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    cur = con.cursor()
    cur.executescript("""
    DROP TABLE IF EXISTS order_profit; DROP TABLE IF EXISTS sales_reps;
    DROP TABLE IF EXISTS customer_value; DROP TABLE IF EXISTS product_profit;
    DROP TABLE IF EXISTS attribution; DROP TABLE IF EXISTS risk_warn;
    CREATE TABLE order_profit (contract_no TEXT, customer TEXT, device TEXT,
        amount REAL, bom_cost REAL, profit REAL, margin REAL);
    CREATE TABLE sales_reps (rep TEXT, order_count INTEGER, sales REAL,
        customer_count INTEGER, est_profit REAL);
    CREATE TABLE customer_value (customer TEXT, total_amount REAL, times INTEGER,
        complaints INTEGER, resp_min REAL, score REAL, tag TEXT);
    CREATE TABLE product_profit (product TEXT, sales REAL, profit REAL);
    CREATE TABLE attribution (contract_no TEXT, customer TEXT, main_rep TEXT,
        tracker TEXT, coworker TEXT, rep_ratio REAL, tracker_ratio REAL,
        coworker_ratio REAL, rep_amount REAL, tracker_amount REAL, coworker_amount REAL);
    CREATE TABLE risk_warn (customer TEXT, amount REAL, times INTEGER,
        reason TEXT, action TEXT);
    CREATE INDEX idx_op_no ON order_profit(contract_no);
    CREATE INDEX idx_at_no ON attribution(contract_no);
    """)

    rows = list(wb["订单利润分析"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO order_profit VALUES (?,?,?,?,?,?,?)",
                    [(s(r[0]), s(r[1]), s(r[2]), f(r[3]), f(r[4]), f(r[5]), f(r[6])) for r in rows if s(r[0])])

    rows = list(wb["销售人员业绩分析"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO sales_reps VALUES (?,?,?,?,?)",
                    [(s(r[0]), int(f(r[1])), f(r[2]), int(f(r[3])), f(r[4])) for r in rows if s(r[0])])

    rows = list(wb["客户价值评估"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO customer_value VALUES (?,?,?,?,?,?,?)",
                    [(s(r[0]), f(r[1]), int(f(r[2])), int(f(r[3])), f(r[4]), f(r[5]), s(r[6])) for r in rows if s(r[0])])

    rows = list(wb["产品利润矩阵"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO product_profit VALUES (?,?,?)",
                    [(s(r[0]), f(r[1]), f(r[2])) for r in rows if s(r[0])])

    rows = list(wb["销售贡献归因"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO attribution VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                    [(s(r[0]), s(r[1]), s(r[2]), s(r[3]), s(r[4]), f(r[5]), f(r[6]),
                      f(r[7]), f(r[8]), f(r[9]), f(r[10])) for r in rows if s(r[0])])

    rows = list(wb["客户风险预警"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO risk_warn VALUES (?,?,?,?,?)",
                    [(s(r[0]), f(r[1]), int(f(r[2])), s(r[3]), s(r[4])) for r in rows if s(r[0])])

    con.commit()
    for t in ("order_profit", "sales_reps", "customer_value", "product_profit",
              "attribution", "risk_warn"):
        print(t, cur.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0])
    con.close()
    print("DB ready:", DB_PATH)


if __name__ == "__main__":
    main()
