# -*- coding: utf-8 -*-
"""金汉隆财务智能化中枢 — 数据导入：xlsx -> SQLite"""
import os
import sqlite3
import sys
from pathlib import Path

import openpyxl

DB_PATH = os.environ.get("FINANCE_DB", os.path.join(os.path.dirname(__file__), "data", "finance.db"))


def f(v, default=0.0):
    try:
        if v is None or v == "":
            return default
        return float(str(v).replace(",", "").replace("¥", "").strip())
    except (ValueError, TypeError):
        return default


def s(v):
    return "" if v is None else str(v).strip()


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "data", "model.xlsx")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    cur = con.cursor()
    cur.executescript("""
    DROP TABLE IF EXISTS reimburse; DROP TABLE IF EXISTS ar; DROP TABLE IF EXISTS ap;
    DROP TABLE IF EXISTS monthly; DROP TABLE IF EXISTS bom; DROP TABLE IF EXISTS sensitivity;
    DROP TABLE IF EXISTS parts; DROP TABLE IF EXISTS orders; DROP TABLE IF EXISTS order_bom;
    CREATE TABLE reimburse (id TEXT, req_date TEXT, fee_type TEXT, title TEXT, tax_no TEXT,
        amount REAL, company TEXT, invoice_date TEXT, forbidden TEXT, result TEXT, reason TEXT);
    CREATE TABLE ar (contract_no TEXT, customer TEXT, occur_date TEXT, amount REAL,
        kind TEXT, aging_days INTEGER, warn60 TEXT, status TEXT);
    CREATE TABLE ap (pay_no TEXT, supplier TEXT, pay_date TEXT, amount REAL, kind TEXT, status TEXT);
    CREATE TABLE monthly (month TEXT, revenue REAL, cost REAL, profit REAL, margin REAL, anomalies INTEGER);
    CREATE TABLE bom (product TEXT, category TEXT, share_pct REAL, cur_cost REAL);
    CREATE TABLE sensitivity (product TEXT, category TEXT, hike_pct REAL, orig_cost REAL,
        pred_cost REAL, cost_delta REAL);
    CREATE TABLE parts (machine TEXT, module TEXT, code TEXT, name TEXT, qty REAL,
        unit TEXT, price REAL, amount REAL, contract_no TEXT, deliver_date TEXT);
    CREATE TABLE orders (order_no TEXT, customer TEXT, device_model TEXT,
        bom_cost REAL, material_count INTEGER);
    CREATE TABLE order_bom (order_no TEXT, seq INTEGER, material_code TEXT, prefix TEXT,
        category TEXT, name TEXT, qty REAL, unit_cost REAL, cost REAL, supply TEXT);
    CREATE INDEX idx_order_bom_no ON order_bom(order_no);
    """)

    rows = list(wb["智能报销审核数据"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO reimburse VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                    [(s(r[0]), s(r[1]), s(r[2]), s(r[3]), s(r[4]), f(r[5]), s(r[6]), s(r[7]), s(r[8]), s(r[9]), s(r[10])) for r in rows])

    rows = list(wb["应收账款台账"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO ar VALUES (?,?,?,?,?,?,?,?)",
                    [(s(r[0]), s(r[1]), s(r[2]), f(r[3]), s(r[4]), int(f(r[5], -1)), s(r[6]), s(r[7])) for r in rows])

    rows = list(wb["应付账款台账"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO ap VALUES (?,?,?,?,?,?)",
                    [(s(r[0]), s(r[1]), s(r[2]), f(r[3]), s(r[4]), s(r[5])) for r in rows])

    # 月度：同月多行聚合
    agg = {}
    rows = list(wb["财务趋势分析看板"].iter_rows(min_row=2, values_only=True))
    for r in rows:
        m = s(r[0])
        if not m:
            continue
        a = agg.setdefault(m, [0.0, 0.0, 0.0, 0])
        a[0] += f(r[1]); a[1] += f(r[2]); a[2] += f(r[3]); a[3] += int(f(r[5]))
    for m, (rev, cost, profit, anom) in sorted(agg.items()):
        margin = profit / rev * 100 if rev else 0
        cur.execute("INSERT INTO monthly VALUES (?,?,?,?,?,?)", (m, rev, cost, profit, round(margin, 2), anom))

    rows = list(wb["产品BOM成本结构"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO bom VALUES (?,?,?,?)",
                    [(s(r[0]), s(r[1]), f(r[2]), f(r[3])) for r in rows])

    rows = list(wb["成本敏感性预测模型"].iter_rows(min_row=2, values_only=True))
    cur.executemany("INSERT INTO sensitivity VALUES (?,?,?,?,?,?)",
                    [(s(r[0]), s(r[1]), f(r[2]), f(r[3]), f(r[4]), f(r[5])) for r in rows])

    if "零件对账明细" in wb.sheetnames:
        rows = list(wb["零件对账明细"].iter_rows(min_row=2, values_only=True))
        cur.executemany("INSERT INTO parts VALUES (?,?,?,?,?,?,?,?,?,?)",
                        [(s(r[0]), s(r[1]), s(r[2]), s(r[3]), f(r[4]), s(r[5]),
                          f(r[6]), f(r[7]), s(r[8]), s(r[9])) for r in rows if s(r[2])])

    # 订单 BOM（CSV）
    import csv as _csv
    for csv_name, sql, ncol in [
        ("orders.csv", "INSERT INTO orders VALUES (?,?,?,?,?)", 5),
        ("order_bom.csv", "INSERT INTO order_bom VALUES (?,?,?,?,?,?,?,?,?,?)", 10),
    ]:
        p = Path(DB_PATH).parent / csv_name
        if p.exists():
            with open(p, encoding="utf-8-sig", newline="") as fh:
                rd = list(_csv.reader(fh))
            cur.executemany(sql, [r[:ncol] for r in rd[1:] if len(r) >= ncol])

    con.commit()
    for t in ("reimburse", "ar", "ap", "monthly", "bom", "sensitivity", "parts",
              "orders", "order_bom"):
        print(t, cur.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0])
    con.close()
    print("DB ready:", DB_PATH)


if __name__ == "__main__":
    main()
