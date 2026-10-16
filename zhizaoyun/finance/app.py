# -*- coding: utf-8 -*-
"""金汉隆 · 财务智能化中枢 — 财务流程自动化
模块：智能报销审核引擎 / 财务数据分析看板 / 智能成本预测模型
基于 datie-bom 骨架改造（Flask + SQLite + Vue2 + ElementUI + ECharts）
"""
import base64
import csv
import io
import json
import os
import sqlite3
from datetime import datetime
from pathlib import Path

import openpyxl
import requests
from flask import Flask, g, jsonify, redirect, render_template, request, send_file

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("FINANCE_DB", BASE_DIR / "data" / "finance.db"))

LLM_BASE = os.environ.get("FINANCE_LLM_BASE", "http://api.computing.lzwcai.com:3110/v1")
LLM_KEY = os.environ.get("FINANCE_LLM_KEY", "")
LLM_MODEL = os.environ.get("FINANCE_LLM_MODEL", "qwen3.8-flash")
OUR_COMPANY = os.environ.get("FINANCE_COMPANY", "广东金汉隆智能装备科技有限公司")

# 差旅与报销标准（元）。可在环境变量里覆盖：FINANCE_STANDARDS={"住宿费":600,...}
DEFAULT_STANDARDS = {
    "住宿费": 600,
    "交通费": 1500,
    "餐饮费": 300,
    "办公费": 2000,
    "通讯费": 500,
}
STANDARDS = json.loads(os.environ.get("FINANCE_STANDARDS", "") or "{}") or DEFAULT_STANDARDS
FORBIDDEN_ITEMS = ["聚餐", "宴请", "招待", "烟", "酒"]
OVERDUE_DAYS = 90          # 发票距报销日超期天数
MANUAL_REVIEW_LIMIT = 10000  # 超过该金额转人工复核

app = Flask(__name__)


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    conn = g.pop("db", None)
    if conn is not None:
        conn.close()


# ----------------------------------------------------------------------------
# 页面
# ----------------------------------------------------------------------------
@app.route("/")
def home():
    return render_template("home.html")


@app.route("/audit")
def audit_page():
    return render_template("audit.html")


@app.route("/dashboard")
def dashboard_page():
    return render_template("dashboard.html")


@app.route("/predict")
def predict_legacy():
    return redirect("/bom", code=302)


@app.route("/bom")
def bom_page():
    return render_template("bom.html")


@app.route("/data")
def data_page():
    return render_template("data.html")


# ----------------------------------------------------------------------------
# 模块一：智能报销审核引擎
# ----------------------------------------------------------------------------
def pdf_to_images(pdf_bytes, max_pages=2, dpi=150):
    """PDF 前几页转 PNG（发票通常一页，保险起见取前两页）"""
    import fitz  # PyMuPDF
    imgs = []
    with fitz.open(stream=pdf_bytes, filetype="pdf") as doc:
        for i, page in enumerate(doc):
            if i >= max_pages:
                break
            pix = page.get_pixmap(dpi=dpi)
            imgs.append(pix.tobytes("png"))
    return imgs


def shrink_image(img_bytes, max_edge=1600):
    """兜底压缩：长边超限的图片缩小并转 JPEG，控制发给视觉模型的体积"""
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(img_bytes))
        if max(im.size) <= max_edge:
            return img_bytes
        im.thumbnail((max_edge, max_edge))
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=85)
        return buf.getvalue()
    except Exception:  # noqa: BLE001
        return img_bytes


def llm_ocr(file_bytes, mime):
    """调用视觉大模型识别发票要素；PDF 自动转图片后识别"""
    if file_bytes[:4] == b"%PDF" or mime == "application/pdf":
        images = pdf_to_images(file_bytes)
        if not images:
            raise ValueError("PDF 无可识别页面")
    else:
        images = [shrink_image(file_bytes)]

    content_parts = [{"type": "text",
                      "text": ("这是发票/报销凭证（可能多张图）。请识别以下要素，只输出 JSON，"
                               "不要输出其他内容：{\"title\":\"发票抬头单位\",\"tax_no\":\"税号\","
                               "\"invoice_date\":\"YYYY-MM-DD\",\"amount\":数字去掉符号(多张取合计),"
                               "\"item\":\"费用项目/商品名称\",\"remark\":\"备注\"}。"
                               "无法识别的字段填空字符串。")}]
    for img in images:
        b64 = base64.b64encode(img).decode()
        content_parts.append({"type": "image_url",
                              "image_url": {"url": f"data:image/png;base64,{b64}"}})

    body = {
        "model": LLM_MODEL,
        "messages": [{"role": "user", "content": content_parts}],
        "max_tokens": 400,
    }
    r = requests.post(LLM_BASE + "/chat/completions",
                      headers={"Authorization": "Bearer " + LLM_KEY},
                      json=body, timeout=90)
    r.raise_for_status()
    content = r.json()["choices"][0]["message"]["content"]
    start, end = content.find("{"), content.rfind("}")
    fields = json.loads(content[start:end + 1]) if start >= 0 else {}
    return {k: fields.get(k, "") for k in
            ("title", "tax_no", "invoice_date", "amount", "item", "remark")}


def audit_rules(fields):
    """规则引擎：差旅标准比对 + 禁项识别 + 超期/抬头校验 → 通过/驳回"""
    reasons, warnings = [], []
    item = str(fields.get("item") or "")
    title = str(fields.get("title") or "")
    try:
        amount = float(str(fields.get("amount") or "0").replace(",", ""))
    except ValueError:
        amount = 0.0

    # 1) 禁项
    hit = [w for w in FORBIDDEN_ITEMS if w in item or w in str(fields.get("remark") or "")]
    if hit:
        reasons.append(f"禁项识别：费用项目含「{ '」「'.join(hit) }」，公司规定不予报销")

    # 2) 抬头校验
    if title and OUR_COMPANY not in title:
        reasons.append(f"发票抬头（{title}）与本公司（{OUR_COMPANY}）不一致")

    # 3) 超期
    inv_date = str(fields.get("invoice_date") or "")
    if inv_date[:10].count("-") == 2:
        try:
            delta = (datetime.now() - datetime.strptime(inv_date[:10], "%Y-%m-%d")).days
            if delta > OVERDUE_DAYS:
                reasons.append(f"发票日期距今日 {delta} 天，超过 {OVERDUE_DAYS} 天报销期限")
            elif delta < 0:
                reasons.append("发票日期晚于今天，疑似信息有误")
        except ValueError:
            warnings.append("发票日期格式无法解析，请人工核对")
    else:
        warnings.append("未识别到发票日期")

    # 4) 金额标准
    std_key = next((k for k in STANDARDS if k and k in item), None)
    if std_key and amount > 0:
        if amount > STANDARDS[std_key]:
            reasons.append(f"{std_key}金额 {amount:.2f} 元超过标准 {STANDARDS[std_key]} 元")
        else:
            warnings.append(f"{std_key} {amount:.2f} 元，标准内（≤{STANDARDS[std_key]} 元）")
    if amount > MANUAL_REVIEW_LIMIT:
        warnings.append(f"金额超过 {MANUAL_REVIEW_LIMIT} 元，需人工复核")

    result = "驳回" if reasons else "通过"
    return {"result": result, "reasons": reasons, "warnings": warnings,
            "amount": amount, "audited_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}


@app.route("/api/audit/ocr", methods=["POST"])
def api_ocr():
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "缺少文件"}), 400
    try:
        fields = llm_ocr(f.read(), f.mimetype or "image/jpeg")
    except Exception as e:  # noqa: BLE001
        return jsonify({"error": f"识别失败：{e}"}), 502
    return jsonify({"fields": fields})


@app.route("/api/audit/submit", methods=["POST"])
def api_submit():
    fields = request.get_json(force=True)
    verdict = audit_rules(fields)
    no = "BXSQ" + datetime.now().strftime("%H%M%S") + f"{int(datetime.now().timestamp()) % 1000:03d}"
    fee_type = str(fields.get("item") or "其他")[:20]
    forbidden = next((w for w in FORBIDDEN_ITEMS
                      if w in fee_type or w in str(fields.get("remark") or "")), "")
    reason = "；".join(verdict["reasons"]) or "符合报销标准"
    try:
        db().execute(
            "INSERT INTO reimburse VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (no, datetime.now().strftime("%Y-%m-%d"), fee_type,
             str(fields.get("title") or ""), str(fields.get("tax_no") or ""),
             verdict["amount"], OUR_COMPANY,
             str(fields.get("invoice_date") or ""), forbidden or "-",
             verdict["result"], reason))
        db().commit()
    except sqlite3.Error:
        pass
    verdict["req_no"] = no
    verdict["reason"] = reason
    return jsonify(verdict)


@app.route("/api/reimburse")
def api_reimburse():
    page, size = int(request.args.get("page", 1)), int(request.args.get("size", 15))
    q = request.args.get("q", "").strip()
    where, args = "", []
    if q:
        where = "WHERE id LIKE ? OR fee_type LIKE ? OR title LIKE ?"
        args = [f"%{q}%"] * 3
    total = db().execute(f"SELECT COUNT(*) c FROM reimburse {where}", args).fetchone()["c"]
    rows = db().execute(
        f"SELECT * FROM reimburse {where} ORDER BY req_date DESC, rowid DESC LIMIT ? OFFSET ?",
        args + [size, (page - 1) * size]).fetchall()
    stat = db().execute(
        "SELECT result, COUNT(*) c FROM reimburse GROUP BY result").fetchall()
    return jsonify({"total": total, "rows": [dict(r) for r in rows],
                    "stat": {r["result"]: r["c"] for r in stat}})


@app.route("/api/standards")
def api_standards():
    return jsonify({"standards": STANDARDS, "forbidden": FORBIDDEN_ITEMS,
                    "overdue_days": OVERDUE_DAYS, "company": OUR_COMPANY})


# ----------------------------------------------------------------------------
# 模块二：财务数据分析看板
# ----------------------------------------------------------------------------
@app.route("/api/overview")
def api_overview():
    d = db()
    ar_total = d.execute("SELECT COALESCE(SUM(amount),0) v FROM ar").fetchone()["v"]
    ar_over60 = d.execute(
        "SELECT COALESCE(SUM(amount),0) v FROM ar WHERE aging_days > 60").fetchone()["v"]
    ap_total = d.execute(
        "SELECT COALESCE(SUM(amount),0) v FROM ap WHERE status='待付款'").fetchone()["v"]
    m = d.execute("SELECT SUM(revenue) r, SUM(cost) c, SUM(profit) p FROM monthly").fetchone()
    margin = (m["p"] or 0) / m["r"] * 100 if m["r"] else 0
    return jsonify({
        "ar_total": ar_total, "ar_over60": ar_over60,
        "ap_pending": ap_total,
        "revenue": m["r"] or 0, "profit": m["p"] or 0, "margin": round(margin, 2),
        "ar_warn_count": d.execute(
            "SELECT COUNT(*) c FROM ar WHERE aging_days > 60").fetchone()["c"],
    })


@app.route("/api/monthly")
def api_monthly():
    rows = db().execute("SELECT * FROM monthly ORDER BY month").fetchall()
    return jsonify({"rows": [dict(r) for r in rows]})


@app.route("/api/ar")
def api_ar():
    d = db()
    buckets = {"0-30天": 0, "31-60天": 0, "61-90天": 0, "90天以上": 0}
    for r in d.execute("SELECT aging_days a, SUM(amount) v FROM ar GROUP BY aging_days"):
        a = r["a"] or 0
        k = "0-30天" if a <= 30 else "31-60天" if a <= 60 else "61-90天" if a <= 90 else "90天以上"
        buckets[k] += r["v"]
    top = d.execute(
        "SELECT customer, SUM(amount) v, MAX(aging_days) max_age, COUNT(*) n "
        "FROM ar GROUP BY customer ORDER BY v DESC LIMIT 10").fetchall()
    kinds = d.execute(
        "SELECT kind, SUM(amount) v FROM ar GROUP BY kind ORDER BY v DESC").fetchall()
    status = d.execute(
        "SELECT status, COUNT(*) n, SUM(amount) v FROM ar GROUP BY status").fetchall()
    return jsonify({"buckets": buckets, "top": [dict(r) for r in top],
                    "kinds": [dict(r) for r in kinds],
                    "status": [dict(r) for r in status]})


@app.route("/api/ar/list")
def api_ar_list():
    only_warn = request.args.get("warn") == "1"
    sql = "SELECT * FROM ar" + (" WHERE aging_days > 60" if only_warn else "") + \
        " ORDER BY aging_days DESC LIMIT 200"
    return jsonify({"rows": [dict(r) for r in db().execute(sql)]})


@app.route("/api/ap")
def api_ap():
    d = db()
    by_month = d.execute(
        "SELECT substr(pay_date,1,7) m, SUM(amount) v FROM ap "
        "GROUP BY m ORDER BY m").fetchall()
    by_type = d.execute(
        "SELECT kind, SUM(amount) v FROM ap GROUP BY kind ORDER BY v DESC").fetchall()
    by_status = d.execute(
        "SELECT status, COUNT(*) n, SUM(amount) v FROM ap GROUP BY status").fetchall()
    top = d.execute(
        "SELECT supplier, SUM(amount) v, COUNT(*) n FROM ap "
        "GROUP BY supplier ORDER BY v DESC LIMIT 10").fetchall()
    return jsonify({"by_month": [dict(r) for r in by_month],
                    "by_type": [dict(r) for r in by_type],
                    "by_status": [dict(r) for r in by_status],
                    "top": [dict(r) for r in top]})


@app.route("/api/alerts")
def api_alerts():
    """异常预警：账龄超60天应收 + 月度费用/毛利异常"""
    alerts = []
    d = db()
    for r in d.execute(
            "SELECT customer, contract_no, amount, aging_days, status FROM ar "
            "WHERE aging_days > 60 ORDER BY aging_days DESC LIMIT 50"):
        alerts.append({
            "level": "高" if r["aging_days"] > 90 else "中",
            "type": "应收账龄",
            "msg": f"客户「{r['customer']}」合同 {r['contract_no']} "
                   f"账龄 {r['aging_days']} 天，金额 {r['amount']:,.0f} 元（{r['status']}）",
        })
    months = [dict(r) for r in d.execute("SELECT * FROM monthly ORDER BY month")]
    cost_avg = sum(m["cost"] for m in months) / max(len(months), 1)
    prev_margin = None
    for m in months:
        if m["cost"] > cost_avg * 1.15:
            alerts.append({
                "level": "中", "type": "费用超预算",
                "msg": f"{m['month']} 生产成本 {m['cost']:,.0f} 元，高于期间均值 "
                       f"{cost_avg:,.0f} 元 15% 以上",
            })
        if prev_margin is not None and prev_margin - m["margin"] >= 5:
            alerts.append({
                "level": "高", "type": "毛利率下滑",
                "msg": f"{m['month']} 毛利率 {m['margin']:.1f}%，环比下滑 "
                       f"{prev_margin - m['margin']:.1f} 个百分点",
            })
        prev_margin = m["margin"]
    return jsonify({"alerts": alerts, "months": months})


# ----------------------------------------------------------------------------
# 模块三：智能成本预测模型 —— 订单级 BOM（可搜索订单/设备/客户）
# ----------------------------------------------------------------------------
@app.route("/api/orders/stats")
def api_orders_stats():
    d = db()
    agg = d.execute("SELECT COUNT(*) n, SUM(bom_cost) total FROM orders").fetchone()
    models = d.execute(
        "SELECT device_model, COUNT(*) n, SUM(bom_cost) total FROM orders "
        "GROUP BY device_model ORDER BY total DESC LIMIT 10").fetchall()
    return jsonify({"order_count": agg["n"], "total_cost": agg["total"] or 0,
                    "top_models": [dict(r) for r in models]})


@app.route("/api/orders/search")
def api_orders_search():
    """按订单号 / 客户 / 设备型号搜索"""
    q = request.args.get("q", "").strip()
    page = max(int(request.args.get("page", 1)), 1)
    size = min(int(request.args.get("size", 30)), 100)
    where, args = "", []
    if q:
        where = ("WHERE order_no LIKE ? OR customer LIKE ? OR device_model LIKE ?")
        args = [f"%{q}%"] * 3
    total = db().execute(f"SELECT COUNT(*) c FROM orders {where}", args).fetchone()["c"]
    rows = db().execute(
        f"SELECT * FROM orders {where} ORDER BY bom_cost DESC LIMIT ? OFFSET ?",
        args + [size, (page - 1) * size]).fetchall()
    return jsonify({"total": total, "rows": [dict(r) for r in rows]})


@app.route("/api/orders/<order_no>/tree")
def api_order_tree(order_no):
    """订单 BOM 成本树：订单 → 物料类别 → 图号组 → 物料"""
    order = db().execute(
        "SELECT * FROM orders WHERE order_no = ?", (order_no,)).fetchone()
    if not order:
        return jsonify({"error": "未找到该订单"}), 404
    parts = db().execute(
        "SELECT order_no, seq, material_code, prefix, category, name, qty, "
        "unit_cost, cost, supply FROM order_bom WHERE order_no = ? ORDER BY cost DESC",
        (order_no,)).fetchall()
    if not parts:
        return jsonify({"error": "该订单没有物料明细"}), 404

    cats = {}
    for p in parts:
        c = cats.setdefault(p["category"] or "其他", {"parts": [], "total": 0.0})
        c["parts"].append(p)
        c["total"] += p["cost"] or 0

    cat_nodes = []
    for cname, c in sorted(cats.items(), key=lambda x: -x[1]["total"]):
        prefixes = {}
        for p in c["parts"]:
            g = prefixes.setdefault(p["prefix"] or "其他", {"parts": [], "total": 0.0})
            g["parts"].append(p)
            g["total"] += p["cost"] or 0
        prefix_nodes = [{
            "name": gname,
            "value": round(g["total"], 2), "totalCost": round(g["total"], 2),
            "usageQty": f"{len(g['parts'])}种物料", "materialName": gname,
            "nodeType": "group", "symbolSize": 20,
            "itemStyle": {"color": "#409eff"},
            "children": [{
                "name": p["name"] or p["material_code"],
                "value": round(p["cost"] or 0, 2),
                "code": p["material_code"], "qty": p["qty"],
                "price": p["unit_cost"], "amount": round(p["cost"] or 0, 2),
                "totalCost": round(p["cost"] or 0, 2),
                "usageQty": f"{p['qty'] or 0}件", "materialName": p["name"],
                "supply": p["supply"] or "",
                "category": cname, "nodeType": "part", "symbolSize": 10,
            } for p in sorted(g["parts"], key=lambda x: -(x["cost"] or 0))],
        } for gname, g in sorted(prefixes.items(), key=lambda x: -x[1]["total"])]
        cat_nodes.append({
            "name": cname,
            "value": round(c["total"], 2), "totalCost": round(c["total"], 2),
            "usageQty": f"{len(c['parts'])}种物料", "materialName": cname,
            "nodeType": "category", "symbolSize": 28,
            "itemStyle": {"color": "#e6a23c"},
            "children": prefix_nodes,
        })

    tree = {
        "name": order_no, "materialName": order_no,
        "value": order["bom_cost"], "totalCost": order["bom_cost"],
        "usageQty": f"{order['material_count']}种物料",
        "customer": order["customer"], "device_model": order["device_model"],
        "nodeType": "root", "symbolSize": 42,
        "itemStyle": {"color": "#2d5b8e"}, "children": cat_nodes,
    }
    return jsonify({"tree": tree, "total": order["bom_cost"],
                    "part_count": order["material_count"]})


# ----------------------------------------------------------------------------
# 数据管理：导入 / 导出 / 模板（导入后看板自动重绘）
# ----------------------------------------------------------------------------
DATA_TABLES = {
    "reimburse": {
        "name": "报销审核数据", "icon": "🧾", "table": "reimburse",
        "headers": ["报销单号", "日期", "费用类型", "发票抬头单位", "税号", "金额",
                    "公司名称", "发票日期", "禁项识别", "审核结果", "审核原因"],
        "cols": ["id", "req_date", "fee_type", "title", "tax_no", "amount",
                 "company", "invoice_date", "forbidden", "result", "reason"],
        "numeric": {"amount"},
    },
    "ar": {
        "name": "应收账款台账", "icon": "📥", "table": "ar",
        "headers": ["合同编号", "客户", "发生日期", "金额", "款项类型", "账龄天数", "60天预警", "状态"],
        "cols": ["contract_no", "customer", "occur_date", "amount", "kind",
                 "aging_days", "warn60", "status"],
        "numeric": {"amount", "aging_days"},
    },
    "ap": {
        "name": "应付账款台账", "icon": "📤", "table": "ap",
        "headers": ["付款编号", "供应商", "日期", "金额", "类型", "状态"],
        "cols": ["pay_no", "supplier", "pay_date", "amount", "kind", "status"],
        "numeric": {"amount"},
    },
    "monthly": {
        "name": "月度经营数据", "icon": "📅", "table": "monthly",
        "headers": ["月份", "销售收入", "生产成本", "毛利润", "毛利率", "异常数量"],
        "cols": ["month", "revenue", "cost", "profit", "margin", "anomalies"],
        "numeric": {"revenue", "cost", "profit", "margin", "anomalies"},
    },
    "parts": {
        "name": "零件对账明细(BOM)", "icon": "⚙️", "table": "parts",
        "headers": ["机器型号", "模块", "图号", "名称", "数量", "单位", "单价", "金额", "合同编号", "送货日期"],
        "cols": ["machine", "module", "code", "name", "qty", "unit",
                 "price", "amount", "contract_no", "deliver_date"],
        "numeric": {"qty", "price", "amount"},
    },
}


def _read_sheet(file_storage):
    """解析上传的 xlsx/csv → (headers, rows)；按表头名映射，忽略多余列"""
    filename = (file_storage.filename or "").lower()
    if filename.endswith(".csv"):
        stream = io.StringIO(file_storage.read().decode("utf-8-sig"))
        data = list(csv.reader(stream))
    else:
        wb = openpyxl.load_workbook(file_storage, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        data = [[("" if c is None else str(c).strip()) for c in r]
                for r in ws.iter_rows(values_only=True)]
    if not data:
        raise ValueError("文件为空")
    return data[0], data[1:]


def _rows_to_tuples(headers, rows, spec):
    idx = {h.strip(): i for i, h in enumerate(headers)}
    missing = [h for h in spec["headers"] if h not in idx]
    if missing:
        raise ValueError(f"缺少列：{'、'.join(missing)}（请使用模板）")
    out = []
    for r in rows:
        if not any(str(x).strip() for x in r):
            continue
        tup = []
        for col, header in zip(spec["cols"], spec["headers"]):
            v = r[idx[header]]
            if col in spec["numeric"]:
                try:
                    v = float(str(v).replace(",", "") or 0)
                    v = int(v) if col in ("aging_days", "anomalies") else v
                except ValueError:
                    v = 0
            tup.append(v)
        out.append(tup)
    return out


def _xlsx_response(spec, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec["name"]
    ws.append(spec["headers"])
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{spec['name']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/api/data/tables")
def api_data_tables():
    out = []
    for key, spec in DATA_TABLES.items():
        n = db().execute(f"SELECT COUNT(*) c FROM {spec['table']}").fetchone()["c"]
        out.append({"key": key, "name": spec["name"], "icon": spec["icon"],
                    "headers": spec["headers"], "count": n})
    return jsonify({"tables": out})


@app.route("/api/data/export/<key>")
def api_data_export(key):
    spec = DATA_TABLES.get(key)
    if not spec:
        return jsonify({"error": "未知数据表"}), 404
    rows = db().execute(f"SELECT {', '.join(spec['cols'])} FROM {spec['table']}").fetchall()
    return _xlsx_response(spec, [tuple(r) for r in rows])


@app.route("/api/data/template/<key>")
def api_data_template(key):
    spec = DATA_TABLES.get(key)
    if not spec:
        return jsonify({"error": "未知数据表"}), 404
    return _xlsx_response(spec, [])


@app.route("/api/data/import/<key>", methods=["POST"])
def api_data_import(key):
    spec = DATA_TABLES.get(key)
    if not spec:
        return jsonify({"error": "未知数据表"}), 404
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "缺少文件"}), 400
    try:
        headers, raw_rows = _read_sheet(f)
        tuples = _rows_to_tuples(headers, raw_rows, spec)
    except Exception as e:  # noqa: BLE001
        return jsonify({"error": f"解析失败：{e}"}), 400
    if not tuples:
        return jsonify({"error": "没有可导入的数据行"}), 400

    conn = db()
    placeholders = ",".join("?" * len(spec["cols"]))
    cols = ",".join(spec["cols"])
    try:
        conn.execute(f"DELETE FROM {spec['table']}")
        conn.executemany(f"INSERT INTO {spec['table']} ({cols}) VALUES ({placeholders})", tuples)
        conn.commit()
    except sqlite3.Error as e:
        conn.rollback()
        return jsonify({"error": f"写入失败：{e}"}), 500
    return jsonify({"imported": len(tuples), "table": spec["name"]})


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5050)), debug=False)
