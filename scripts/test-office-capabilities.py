"""Real file-path acceptance checks. Run with the project's Python, never global Python.

All data is synthetic and explicitly labelled. No model, account or customer data is used.
Optional services are reported as untested, never as passed. Outputs stay outside git.
"""
import argparse
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
import traceback


def run(output):
    output.mkdir(parents=True, exist_ok=True)
    results = []

    def check(name, action):
        start = time.perf_counter()
        try:
            detail = action()
            results.append(dict(name=name, status="pass", seconds=round(time.perf_counter()-start, 3), detail=detail))
        except Exception as exc:
            results.append(dict(name=name, status="fail", seconds=round(time.perf_counter()-start, 3), detail=str(exc)))
            traceback.print_exc()

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches as DocxInches
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, Reference
    from pptx import Presentation
    from pptx.util import Inches, Pt as SlidePt
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from PIL import Image, ImageDraw, ImageFont
    from fpdf import FPDF
    from pypdf import PdfReader, PdfWriter
    import pdfplumber
    import pypdfium2 as pdfium

    label = "办公能力测试（合成数据）"
    font = next((p for p in [Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simhei.ttf"),
                              Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")] if p.exists()), None)
    image = output / "测试素材.png"
    im = Image.new("RGB", (1400, 500), "white")
    draw = ImageDraw.Draw(im)
    draw.text((60, 60), "OFFICE TEST 2026", font=ImageFont.truetype(str(font), 80) if font else None, fill="black")
    if font and os.name == "nt":
        draw.text((60, 210), "办公测试 合计 85.00", font=ImageFont.truetype(str(font), 72), fill="black")
    im.save(image)

    def word():
        path = output / "测试 #合同%20_v1.docx"
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(11)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        doc.add_heading(label, 0)
        doc.add_heading("一、验收范围", 1)
        doc.add_paragraph("中文段落、样式、表格、图片与页码字段。")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "项目", "金额"
        cells = table.add_row().cells
        cells[0].text, cells[1].text = "合成样本", "123.45"
        doc.add_picture(str(image), width=DocxInches(4))
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        doc.sections[0].footer.paragraphs[0]._p.append(field)
        doc.save(path)
        original = hashlib.sha256(path.read_bytes()).hexdigest()
        reopened = Document(path)
        assert reopened.tables[0].cell(1, 1).text == "123.45"
        assert len(reopened.inline_shapes) == 1
        assert reopened.paragraphs[1].style.name == "Heading 1"
        reopened.paragraphs[2].add_run(" 已编辑。")
        edited = output / "测试合同_v2.docx"
        reopened.save(edited)
        assert "已编辑" in Document(edited).paragraphs[2].text
        assert hashlib.sha256(path.read_bytes()).hexdigest() == original
        return "中文、标题样式、表格、图片、页码域、编辑重读和原件完整性通过；视觉排版需渲染"

    def excel():
        wb = Workbook()
        ws = wb.active
        ws.title = "合成明细"
        ws.append(["项目", "数量", "单价", "金额"])
        ws.append(["合成A", 2, 12.5, "=B2*C2"])
        ws.append(["合成B", 3, 20, "=B3*C3"])
        ws.append(["合计", None, None, "=SUM(D2:D3)"])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:D3"
        ws.print_title_rows = "1:1"
        for row in ws.iter_rows(min_row=2, max_row=4, min_col=3, max_col=4):
            for cell in row: cell.number_format = '#,##0.00'
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, max_col=2, min_row=1, max_row=3), titles_from_data=True)
        ws.add_chart(chart, "F2")
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        path = output / "测试 #报表%20.xlsx"
        wb.save(path)
        original = hashlib.sha256(path.read_bytes()).hexdigest()
        reopened = load_workbook(path)
        assert reopened.active["D4"].value == "=SUM(D2:D3)"
        assert reopened.active.freeze_panes == "A2" and len(reopened.active._charts) == 1
        assert reopened.calculation.fullCalcOnLoad
        reopened.active["B2"] = 4
        reopened.save(output / "测试报表_v2.xlsx")
        assert load_workbook(output / "测试报表_v2.xlsx").active["D2"].value == "=B2*C2"
        assert hashlib.sha256(path.read_bytes()).hexdigest() == original
        # CSV GB18030 -> UTF-8 BOM; preserve leading zeros as text.
        raw = "编号,金额\n001,25\n002,60\n".encode("gb18030")
        rows = list(csv.reader(io.StringIO(raw.decode("gb18030"))))
        assert rows[1][0] == "001" and sum(float(row[1]) for row in rows[1:]) == 85
        with (output / "中文CSV.csv").open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
        assert (output / "中文CSV.csv").read_bytes().startswith(b'\xef\xbb\xbf')
        return "工作表/公式持久化/格式/冻结/图表/打印/中文CSV通过；公式缓存由单独重算检查验证"

    def ppt():
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = label
        slide.shapes.title.text_frame.paragraphs[0].font.size = SlidePt(28)
        slide.shapes.add_picture(str(image), Inches(1), Inches(1.5), width=Inches(5))
        data = ChartData()
        data.categories = ["合成A", "合成B"]
        data.add_series("样本数量", (2, 3))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(2), Inches(5), Inches(4), data)
        slide.notes_slide.notes_text_frame.text = "仅用于功能测试，不代表业务数据。"
        path = output / "测试 #汇报%20.pptx"
        deck.save(path)
        reopened = Presentation(path)
        assert len(reopened.slides) == 1
        assert any(s.has_chart for s in reopened.slides[0].shapes)
        assert label == reopened.slides[0].shapes.title.text
        assert "功能测试" in reopened.slides[0].notes_slide.notes_text_frame.text
        for shape in reopened.slides[0].shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left+shape.width <= reopened.slide_width+1000
            assert shape.top+shape.height <= reopened.slide_height+1000
        reopened.slides[0].shapes.title.text = label + " 已编辑"
        reopened.slides[0].shapes.title.text_frame.paragraphs[0].font.size = SlidePt(28)
        reopened.save(output / "测试汇报_v2.pptx")
        return "16:9/中文/图片/原生图表/备注/元素边界/编辑重读通过；文本视觉溢出另需渲染"

    def pdf():
        doc = FPDF()
        if font:
            doc.add_font("Test", fname=str(font))
            doc.set_font("Test", size=16)
        else:
            doc.set_font("Helvetica", size=16)
        doc.add_page()
        doc.cell(0, 12, "OFFICE TEST 2026", new_x="LMARGIN", new_y="NEXT")
        if font and os.name == "nt": doc.cell(0, 12, label, new_x="LMARGIN", new_y="NEXT")
        for a, b in [("Item", "Amount"), ("Synthetic A", "25.00"), ("Synthetic B", "60.00")]:
            doc.cell(85, 12, a, border=1)
            doc.cell(85, 12, b, border=1, new_x="LMARGIN", new_y="NEXT")
        doc.add_page()
        doc.cell(0, 12, "SECOND PAGE")
        path = output / "测试 #文档%20.pdf"
        doc.output(str(path))
        with pdfplumber.open(path) as read:
            assert len(read.pages) == 2 and "OFFICE TEST" in read.pages[0].extract_text()
            tables = read.pages[0].extract_tables()
            assert tables and tables[0][1][1] == "25.00"
            rebuilt = Document()
            for page_index, page in enumerate(read.pages, 1):
                rebuilt.add_heading(f"PDF 第 {page_index} 页（内容重建）", 1)
                rebuilt.add_paragraph(page.extract_text() or "")
            rebuilt.save(output / "PDF内容重建.docx")
        assert "OFFICE TEST" in "\n".join(p.text for p in Document(output / "PDF内容重建.docx").paragraphs)
        split = PdfWriter()
        split.add_page(PdfReader(path).pages[0])
        split.write(output / "拆分.pdf")
        merged = PdfWriter()
        merged.append(output / "拆分.pdf")
        merged.append(path)
        merged.pages[0].rotate(90)
        merged.write(output / "合并旋转.pdf")
        assert len(PdfReader(output / "合并旋转.pdf").pages) == 3
        assert PdfReader(output / "合并旋转.pdf").pages[0].rotation == 90
        encrypted = PdfWriter()
        encrypted.append(path)
        encrypted.encrypt("synthetic-test-password")
        encrypted.write(output / "加密.pdf")
        read = PdfReader(output / "加密.pdf")
        assert read.is_encrypted and read.decrypt("synthetic-test-password")
        assert len(read.pages) == 2
        watermark = FPDF()
        watermark.add_page()
        watermark.set_font("Helvetica", size=24)
        watermark.cell(0, 20, "TEST WATERMARK")
        watermark.output(str(output / "水印模板.pdf"))
        stamped = PdfWriter()
        stamped.append(path)
        stamped.pages[0].merge_page(PdfReader(output / "水印模板.pdf").pages[0])
        stamped.write(output / "加水印.pdf")
        assert "TEST WATERMARK" in PdfReader(output / "加水印.pdf").pages[0].extract_text()
        with pdfium.PdfDocument(path) as rendered:
            page = rendered[0]
            bitmap = page.render(scale=2)
            bitmap.to_pil().save(output / "PDF预览.png")
            bitmap.close()
            page.close()
        assert Image.open(output / "PDF预览.png").width > 1000
        # Image-only PDF actually routes to OCR.
        scan = FPDF()
        scan.add_page()
        scan.image(str(image), x=10, y=10, w=190)
        scan.output(str(output / "扫描样本.pdf"))
        assert not (PdfReader(output / "扫描样本.pdf").pages[0].extract_text() or '').strip()
        return "生成/中文嵌字/表格提取/拆分/合并/旋转/水印/加解密/渲染/PDF内容重建Word/扫描件分流通过"

    def ocr():
        from rapidocr_onnxruntime import RapidOCR
        engine = RapidOCR()
        with pdfium.PdfDocument(output / "扫描样本.pdf") as doc:
            page = doc[0]
            bitmap = page.render(scale=200/72)
            bitmap.to_pil().save(output / "扫描渲染.png")
            bitmap.close()
            page.close()
        lines, _ = engine(str(output / "扫描渲染.png"))
        text = " ".join(row[1] for row in lines or [])
        assert "OFFICE" in text.upper() and "2026" in text, text
        if os.name == "nt": assert "办公" in text and "85.00" in text, text
        return dict(text=text, line_count=len(lines), route="本地扫描PDF→整页渲染→RapidOCR；中英文和金额清晰样本")

    def data_analysis():
        import pandas as pd
        import chardet
        from markdownify import markdownify
        frame = pd.DataFrame({"group": ["A", "A", "B"], "amount": [25, 60, 10]})
        assert frame.groupby("group").amount.sum().to_dict() == {"A": 85, "B": 10}
        text = "中文编码测试" * 20
        raw = text.encode("gb18030")
        encoding = chardet.detect(raw)["encoding"]
        assert raw.decode(encoding) == text
        assert "标题" in markdownify("<h1>标题</h1><p>正文</p>")
        copied = output / "文件归档" / image.name
        copied.parent.mkdir(exist_ok=True)
        shutil.copy2(image, copied)
        assert copied.read_bytes() == image.read_bytes()
        return "pandas分组汇总/中文编码探测/HTML转Markdown/文件副本归档/图片校验通过"

    for name, action in [("word", word), ("excel", excel), ("ppt", ppt), ("pdf", pdf), ("ocr", ocr), ("file-data-image", data_analysis)]:
        check(name, action)

    soffice = os.environ.get("SOFFICE") or shutil.which("soffice") or next((str(p) for p in [Path("C:/Program Files/LibreOffice/program/soffice.exe")] if p.exists()), None)
    if soffice:
        def convert():
            converted = output / "converted"
            converted.mkdir(exist_ok=True)
            profile = (output / "lo-profile").resolve().as_uri()
            for pattern in ["*_v1.docx", "*.pptx"]:
                for source in output.glob(pattern):
                    subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", str(converted), str(source)], check=True, timeout=120, capture_output=True)
                    assert len(PdfReader(converted / (source.stem+".pdf")).pages) >= 1
                    with pdfium.PdfDocument(converted / (source.stem+".pdf")) as rendered:
                        for i in range(len(rendered)):
                            page = rendered[i]
                            bitmap = page.render(scale=1.5)
                            bitmap.to_pil().save(converted / f"{source.stem}-page-{i+1}.png")
                            bitmap.close()
                            page.close()
            source = output / "测试 #报表%20.xlsx"
            subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "xlsx", "--outdir", str(converted), str(source)], check=True, timeout=120, capture_output=True)
            result = load_workbook(converted/source.name, data_only=True)
            assert result.active["D4"].value == 85
            assert load_workbook(converted/source.name).active["D4"].data_type == "f"
            return "Word/PPT导出PDF及Excel公式缓存重算通过"
        check("office-convert-recalculate", convert)
    else:
        results.append(dict(name="office-convert-recalculate", status="untested", detail="未安装 LibreOffice；未启动用户 Office/WPS，Word/PPT渲染、Excel公式缓存未验收"))

    for name, reason in [
        ("browser", "需运行时 browser 工具实测；Python包存在不代表浏览器可启动"),
        ("research-image-search", "需外网检索/图源授权链验证；本次不以文档存在代替功能通过"),
        ("desktop-gui", "需 Desktop/Computer Use 插件与可操作会话"),
        ("content-agent", "需模型提供商与真实对话验收；本脚本仅验证文件数据路径"),
        ("cloud-ocr", "未配置且未上传文件到 MinerU"),
    ]:
        results.append(dict(name=name, status="untested", detail=reason))
    report = dict(python=sys.version, executable=sys.executable, platform=sys.platform,
                  synthetic_data=True, results=results)
    (output / "results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return int(any(row["status"] == "fail" for row in results))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True, help="New/isolated directory for synthetic test artifacts")
    args = parser.parse_args()
    sys.exit(run(args.output.resolve()))
