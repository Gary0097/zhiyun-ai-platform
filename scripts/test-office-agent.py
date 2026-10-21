"""Model -> native console SSE -> real agent tools -> files -> download acceptance.

Configure a provider through the platform first. This harness accepts no secrets.
Only explicitly labelled synthetic fixtures are sent to the agent.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re
import time
import urllib.parse
import urllib.request
import uuid
from datetime import datetime
import os
import subprocess


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--base-url', default='http://127.0.0.1:8088')
    parser.add_argument('--case', choices=['word', 'excel', 'ppt', 'pdf'], required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--verify-dir', type=Path, help='Recheck an existing generated test directory without another model call')
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[1]
    workspace = repo/'apps/zhizaoyunAIOS/workspace/workspaces/default'
    run = uuid.uuid4().hex[:8]
    target = args.verify_dir.resolve() if args.verify_dir else workspace/'files'/f'office-agent-{run}'
    target.mkdir(parents=True, exist_ok=bool(args.verify_dir))
    if args.verify_dir: run = target.name.removeprefix('office-agent-')
    args.output.mkdir(parents=True, exist_ok=True)
    specs = {
        'word': ('docx', 'Word报告：中文标题“办公合成验收”、二级标题、两行数据表、页码。'),
        'excel': ('xlsx', 'Excel：明细表，列为项目/数量/单价/金额；金额用公式，末行SUM合计；冻结首行，数字格式两位小数。设置打开时自动重算。'),
        'ppt': ('pptx', 'PPT：3页，封面、原生柱状图展示两项金额、结论；16:9、中文，标题28pt、正文至少18pt。'),
        'pdf': ('pdf', 'PDF：中文标题“办公合成验收”、两行数据表和合计，嵌入本机微软雅黑字体。'),
    }
    extension, spec = specs[args.case]
    filename = f'办公 #{args.case}%20.{extension}'
    prompt = f'''这是用户授权的真实办公合成验收。请实际调用办公技能与本机工具生成文件，不要只给代码或假链接。
任务：{spec}
所有内容标注“合成测试数据”；数据：合成A数量2、单价12.50、金额25；合成B数量3、单价20、金额60；总计85。
受众是内部测试员，内容与大纲已确认，直接生成，不再提问。读取对应 office-{args.case} 技能，使用项目现成Python库，不安装包、不联网、不委派其他智能体。
只允许在此新目录写入：{target.as_posix()}。最终文件必须精确命名：{filename}。
完成后实际打开重读确认内容，并用Markdown文件链接交付真实路径。文件名含中文、空格、#和%20是故意的下载验收。
禁止读取或输出密钥、供应商配置、secret目录或用户其他文件。需要临时脚本只能放上述测试目录。'''
    body = {'session_id': f'office-acceptance-{args.case}-{run}', 'user_id': 'office-synthetic-test',
            'input': [{'role': 'user', 'content': [{'type': 'text', 'text': prompt}]}], 'stream': True}
    req = urllib.request.Request(args.base_url+'/api/agents/default/console/chat', data=json.dumps(body).encode(), headers={'Content-Type':'application/json'})
    start = time.monotonic()
    events = []
    if not args.verify_dir:
        with urllib.request.urlopen(req, timeout=900) as response:
            for line in response:
                line = line.decode('utf-8').strip()
                if not line.startswith('data:'): continue
                data = line[5:].strip()
                if data == '[DONE]': continue
                try: event = json.loads(data)
                except ValueError: continue
                events.append(event)
    else:
        trace=args.output/f'{args.case}-{run}-events.json'
        if trace.exists(): events=json.loads(trace.read_text(encoding='utf-8'))
    # Defensive redaction before persisting any raw tool/model trace.
    raw = json.dumps(events, ensure_ascii=False)
    raw = re.sub(r'sk-[A-Za-z0-9_-]{12,}', '[REDACTED]', raw)
    raw = re.sub(r'Bearer\s+[A-Za-z0-9._-]+', 'Bearer [REDACTED]', raw)
    if not args.verify_dir:
        (args.output/f'{args.case}-{run}-events.json').write_text(raw, encoding='utf-8')
    path = target/filename
    result = {'case':args.case, 'run':run, 'seconds':round(time.monotonic()-start,2),
              'events':len(events), 'file_created':path.exists(), 'path':str(path), 'session_id':body['session_id']}
    envelopes=[e for e in events if e.get('object')=='response' and e.get('completed_at')]
    if envelopes:
        result['agent_seconds']=(datetime.fromisoformat(envelopes[-1]['completed_at'])-datetime.fromisoformat(envelopes[-1]['created_at'])).total_seconds()
        calls=[content.get('data',{}) for message in envelopes[-1].get('output',[]) if message.get('type')=='plugin_call' for content in message.get('content',[])]
        result['tools_used']=sorted({call.get('name') for call in calls if call.get('name')})
        result['skills_used']=[]
        for call in calls:
            if call.get('name')!='Skill': continue
            arguments=call.get('arguments',{})
            if isinstance(arguments,str): arguments=json.loads(arguments)
            result['skills_used'].append(arguments.get('skill'))
        assert 'office-'+args.case in result['skills_used'],result['skills_used']
    if args.verify_dir:
        result['verification_only']=True
    if path.exists():
        assert path.stat().st_size > 100
        if extension == 'docx':
            from docx import Document
            doc=Document(path)
            assert '办公合成验收' in '\n'.join(p.text for p in doc.paragraphs)
            all_text='\n'.join(p.text for p in doc.paragraphs)+str([[c.text for c in r.cells] for t in doc.tables for r in t.rows])
            assert doc.tables and '85' in all_text
        elif extension == 'xlsx':
            from openpyxl import load_workbook
            wb=load_workbook(path)
            formulas=[c.value for ws in wb for row in ws for c in row if c.data_type=='f']
            assert len(formulas)>=3 and any('SUM(' in f.upper() for f in formulas), formulas
            assert any(ws.freeze_panes for ws in wb)
        elif extension == 'pptx':
            from pptx import Presentation
            deck=Presentation(path)
            assert len(deck.slides)==3
            assert any(shape.has_chart for slide in deck.slides for shape in slide.shapes)
        else:
            from pypdf import PdfReader
            text='\n'.join(p.extract_text() or '' for p in PdfReader(path).pages)
            assert '办公合成验收' in text and '85' in text
        # Same native download route used by the workspace file browser.
        relative=path.relative_to(workspace).as_posix()
        query=urllib.parse.urlencode({'path':relative,'root':'workspace'})
        try:
            with urllib.request.urlopen(args.base_url+'/api/agents/default/workspace/file-download?'+query, timeout=30) as response:
                downloaded=response.read()
            assert hashlib.sha256(downloaded).digest()==hashlib.sha256(path.read_bytes()).digest()
            result['download_verified']=True
            # Match the branded chat-card helper and its two server-side decodes.
            preview_path=path.as_posix().lstrip('/').replace('%','%2525').replace('#','%23').replace('?','%3F')
            preview_url=args.base_url+'/api/files/preview/'+urllib.parse.quote(preview_path,safe='/:%')
            with urllib.request.urlopen(preview_url,timeout=30) as response:
                preview=response.read()
            assert hashlib.sha256(preview).digest()==hashlib.sha256(path.read_bytes()).digest()
            result['chat_card_download_verified']=True
        except Exception as exc:
            result['download_verified']=False
            result['download_error']=str(exc)
        soffice=os.environ.get('SOFFICE')
        rendered=args.output/f'{args.case}-{run}-rendered'
        rendered.mkdir(exist_ok=True)
        if soffice and extension in ('docx','pptx','xlsx'):
            export='xlsx' if extension=='xlsx' else 'pdf'
            profile=(rendered/'profile').resolve().as_uri()
            subprocess.run([soffice,f'-env:UserInstallation={profile}','--headless','--convert-to',export,
                            '--outdir',str(rendered.resolve()),str(path)],capture_output=True,check=True,timeout=120)
            converted=rendered/(path.stem+'.'+export)
            assert converted.exists()
            result['conversion_verified']=True
            if extension=='xlsx':
                from openpyxl import load_workbook
                calculated=load_workbook(converted,data_only=True)
                values=[calculated[ws.title][c.coordinate].value for ws in wb for row in ws for c in row if c.data_type=='f']
                assert all(value in values for value in [25,60,85]),values
                result['calculated_values']=values
        pdf_path=path if extension=='pdf' else rendered/(path.stem+'.pdf')
        if pdf_path.exists():
            import pypdfium2 as pdfium
            with pdfium.PdfDocument(pdf_path) as document:
                for index in range(len(document)):
                    page=document[index]
                    bitmap=page.render(scale=1.2)
                    bitmap.to_pil().save(rendered/f'page-{index+1}.png')
                    bitmap.close()
                    page.close()
            result['pages_rendered']=len(list(rendered.glob('page-*.png')))
    (args.output/f'{args.case}-{run}-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False),flush=True)
    return 0 if result.get('download_verified') else 1


if __name__=='__main__':
    raise SystemExit(main())
