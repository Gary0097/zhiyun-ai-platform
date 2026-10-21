"""Read only the demo's standalone XLSX builder; do not import its services/data."""
import argparse
import io
import json
from pathlib import Path
import subprocess
from openpyxl import load_workbook

parser = argparse.ArgumentParser()
parser.add_argument("--ref", default="ada051fd223912e8940c729c9804e737fc487d42")
parser.add_argument("--output", type=Path, required=True)
args = parser.parse_args()
root = Path(__file__).resolve().parents[1]
source = subprocess.check_output(["git", "show", f"{args.ref}:zhizaoyun/order-center/excelgen.py"], cwd=root).decode("utf-8")
namespace = {"__name__": "demo_xlsx_acceptance"}
exec(compile(source, "demo/order-center/excelgen.py", "exec"), namespace)
record = {"customer": "合成验收客户", "so_no": "TEST-001", "parsed_json": {"客户名称": "合成验收客户"},
          "adapted_json": {"物料清单": [{"物料名称": "合成A", "数量": 2, "标准单价": 12.5},
                                      {"物料名称": "合成B", "数量": 3, "标准单价": 20}]}}
data = namespace["build"](record).getvalue()
wb = load_workbook(io.BytesIO(data), data_only=True)
assert wb.sheetnames == ["订单信息", "物料明细", "解析说明"]
assert wb["物料明细"]["H6"].value == 85
assert wb["订单信息"]["B12"].value == 85
assert "合成验收客户" in namespace["filename"](record)
args.output.mkdir(parents=True, exist_ok=True)
(args.output/"demo-合成工单.xlsx").write_bytes(data)
print(json.dumps({"status": "pass", "ref": args.ref, "sheets": wb.sheetnames, "total": 85,
                  "scope": "Standalone Excel builder only; no services, customer snapshots or credentials merged"}, ensure_ascii=False))
