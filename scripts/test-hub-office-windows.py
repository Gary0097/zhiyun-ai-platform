"""Actual Hub AppContainer conversion acceptance (administrator required).

The child receives the upstream LocalProcessRuntimeProvisioner environment;
no inherited Office/Python/PATH settings or sandbox bypass is used.
"""
import argparse
from concurrent.futures import ThreadPoolExecutor
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import uuid


def child(other):
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation
    from pypdf import PdfReader

    workspace = Path(os.environ['QWENPAW_WORKING_DIR'])
    office = shutil.which('soffice')
    assert office and Path(office).name.lower() == 'soffice.exe'
    assert 'SOFFICE' not in os.environ, 'Must prove native member PATH discovery'
    # Read denial must remain intact for the other tenant's synthetic document.
    try:
        Path(other).read_text()
    except OSError:
        pass
    else:
        raise AssertionError('Other tenant workspace was readable')
    document = Document()
    document.add_heading('Synthetic Hub conversion 中文', 0)
    document.add_paragraph('Total 85')
    document.save(workspace / '合同 #合成%20.docx')
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[0]).shapes.title.text = 'Synthetic total 85'
    slides.save(workspace / '汇报 #合成%20.pptx')
    book = Workbook()
    book.active.append([25, 60, '=SUM(A1:B1)'])
    book.save(workspace / '报表 #合成%20.xlsx')

    def convert(name, extension):
        source = workspace / name
        target = workspace / ('converted-' + extension + '-' + source.suffix[1:])
        target.mkdir()
        result = subprocess.run([office, '--headless', '--convert-to', extension,
                                 '--outdir', str(target), str(source)],
                                capture_output=True, text=True, timeout=120)
        assert result.returncode == 0, result.stderr
        output = target / (source.stem + '.' + extension)
        assert output.is_file(), result.stdout
        if extension == 'pdf':
            assert '85' in ''.join(page.extract_text() for page in PdfReader(output).pages)
        else:
            assert load_workbook(output, data_only=True).active['C1'].value == 85
        return output.name

    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(convert, '合同 #合成%20.docx', 'pdf'),
                   pool.submit(convert, '汇报 #合成%20.pptx', 'pdf')]
        for future in futures:
            future.result()
    convert('报表 #合成%20.xlsx', 'xlsx')
    assert not list((workspace / 'tmp').glob('aios-office-*'))
    (workspace / 'conversion-result.json').write_text(json.dumps({
        'word_to_pdf': 'pass', 'ppt_to_pdf': 'pass', 'excel_recalculation': 'pass',
        'concurrent_profiles': 'pass', 'other_tenant_denied': 'pass',
        'office_discovered_on_member_path': 'pass'}), encoding='utf-8')


def parent(root, output):
    from qwenpaw.hub.local_provisioner import LocalProcessRuntimeProvisioner
    from qwenpaw.hub.models import RuntimeRecord, RuntimeState
    from qwenpaw.hub.windows_process_isolation import WindowsAppContainerIsolator
    from qwenpaw.utils.platform import is_windows_admin
    assert sys.platform == 'win32' and is_windows_admin(), 'Native Hub AppContainer requires Windows administrator privileges'
    expected = root / 'apps/zhizaoyunAIOS/runtime/qwenpaw-hub/venv'
    assert Path(sys.prefix).resolve() == expected.resolve(), 'Run with packaged Hub Python'
    assert (root / '.synthetic-runtime-acceptance').read_text().strip() == 'synthetic-only'
    output.mkdir(parents=True, exist_ok=False)
    records = []
    for name in ['alice', 'bob']:
        base = output / name
        record = RuntimeRecord(runtime_id='office-' + uuid.uuid4().hex, tenant_id=name,
                               owner_user_id=name, provisioner='local', host='127.0.0.1',
                               port=18990, state=RuntimeState.CREATED,
                               working_dir=base / 'working', secret_dir=base / 'secrets',
                               backup_dir=base / 'backups', log_file=base / 'logs/app.log')
        for folder in [record.working_dir, record.working_dir / 'tmp',
                       record.working_dir / 'appdata/roaming', record.working_dir / 'appdata/local',
                       record.secret_dir, record.backup_dir, record.log_file.parent]:
            folder.mkdir(parents=True, exist_ok=True)
        (record.working_dir / 'private.txt').write_text('Synthetic tenant data')
        shutil.copy2(__file__, record.working_dir / 'test-conversion.py')
        records.append(record)
    for record, other in [(records[0], records[1]), (records[1], records[0])]:
        environment = LocalProcessRuntimeProvisioner.runtime_environment(record, {})
        isolator = WindowsAppContainerIsolator()
        try:
            launch = isolator.prepare(record, [sys.executable, '-B', str(record.working_dir / 'test-conversion.py'),
                                               '--child', str(other.working_dir / 'private.txt')], environment)
            with record.log_file.open('w', encoding='utf-8') as log:
                process = isolator.launch(record, launch, log)
                try:
                    assert process.wait(timeout=240) == 0, f'Conversion failed: {record.log_file}'
                finally:
                    if process.poll() is None:
                        process.kill()
            assert (record.working_dir / 'conversion-result.json').is_file()
        finally:
            isolator.release(record.runtime_id)
    (output / 'PASS.json').write_text(json.dumps({'native_windows_appcontainer': True, 'tenants': 2,
                                                 'conversion_and_isolation': 'pass'}), encoding='utf-8')
    print('PASS two native Hub AppContainer members: conversions, recalculation, concurrent profiles and tenant read denial', flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--root', type=Path)
    parser.add_argument('--output', type=Path)
    parser.add_argument('--child')
    args = parser.parse_args()
    if args.child:
        child(args.child)
    else:
        parent(args.root.resolve(), args.output.resolve())
